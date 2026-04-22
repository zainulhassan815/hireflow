"""Unit tests for F105.c's SpreadsheetProvider.

Generates in-memory xlsx via openpyxl so the fixtures are
deterministic and self-describing.
"""

from __future__ import annotations

import json
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.adapters.protocols import StoredBlob
from app.adapters.viewers import SpreadsheetProvider
from app.adapters.viewers._table_shape import (
    MAX_COLS_PER_SHEET,
    MAX_ROWS_PER_SHEET,
)
from app.models import Document

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _doc(mime: str = XLSX_MIME, *, filename: str = "file.xlsx") -> Document:
    return Document(
        id=uuid4(),
        owner_id=uuid4(),
        filename=filename,
        mime_type=mime,
        size_bytes=1024,
        storage_key=f"test/{filename}",
    )


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    # Strip the default "Sheet" so tests see exactly what they wrote.
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _FakeStorage:
    def __init__(self, initial: dict[str, bytes] | None = None) -> None:
        self.store: dict[str, bytes] = dict(initial or {})

    def get_sync(self, key: str) -> bytes:
        return self.store[key]

    def put_sync(self, key: str, data: bytes, content_type: str) -> StoredBlob:
        self.store[key] = data
        return StoredBlob(key=key, size=len(data), etag="fake")


def test_accepts_only_xlsx() -> None:
    p = SpreadsheetProvider()
    assert p.accepts(XLSX_MIME)
    for mime in (
        "application/vnd.ms-excel",  # .xls — legacy, out of scope
        "text/csv",
        "application/pdf",
        "application/vnd.oasis.opendocument.spreadsheet",  # .ods
        None,
    ):
        assert not p.accepts(mime), mime


def test_prepare_single_sheet_round_trip() -> None:
    source = _xlsx(
        {
            "Employees": [
                ["Name", "Role", "Years"],
                ["Alice", "Engineer", 5],
                ["Bob", "Designer", 3],
            ]
        }
    )
    storage = _FakeStorage({"test/file.xlsx": source})
    doc = _doc()

    result = SpreadsheetProvider().prepare(
        doc, storage_get=storage.get_sync, storage_put=storage.put_sync
    )

    assert result.kind == "table"
    assert result.key == f"viewable/{doc.id}.json"
    payload = json.loads(storage.store[result.key].decode("utf-8"))
    assert len(payload["sheets"]) == 1
    sheet = payload["sheets"][0]
    assert sheet["name"] == "Employees"
    assert sheet["headers"] == ["Name", "Role", "Years"]
    # Cells serialise to strings for a uniform grid.
    assert sheet["rows"] == [
        ["Alice", "Engineer", "5"],
        ["Bob", "Designer", "3"],
    ]
    assert sheet["truncated"] is False
    assert sheet["total_rows"] == 2
    assert sheet["total_cols"] == 3


def test_prepare_multi_sheet_preserves_order() -> None:
    source = _xlsx(
        {
            "Summary": [["Total", "42"]],
            "Detail": [["A", "B"], ["1", "2"]],
        }
    )
    storage = _FakeStorage({"test/file.xlsx": source})
    doc = _doc()

    SpreadsheetProvider().prepare(
        doc, storage_get=storage.get_sync, storage_put=storage.put_sync
    )

    payload = json.loads(storage.store[f"viewable/{doc.id}.json"].decode("utf-8"))
    names = [s["name"] for s in payload["sheets"]]
    assert names == ["Summary", "Detail"]


def test_prepare_row_cap() -> None:
    rows = [["Col"]] + [[str(i)] for i in range(MAX_ROWS_PER_SHEET + 50)]
    source = _xlsx({"Big": rows})
    storage = _FakeStorage({"test/file.xlsx": source})
    doc = _doc()

    SpreadsheetProvider().prepare(
        doc, storage_get=storage.get_sync, storage_put=storage.put_sync
    )

    sheet = json.loads(storage.store[f"viewable/{doc.id}.json"].decode("utf-8"))[
        "sheets"
    ][0]
    assert len(sheet["rows"]) == MAX_ROWS_PER_SHEET
    assert sheet["total_rows"] == MAX_ROWS_PER_SHEET + 50
    assert sheet["truncated"] is True


def test_prepare_col_cap() -> None:
    wide_row = [f"c{i}" for i in range(MAX_COLS_PER_SHEET + 10)]
    source = _xlsx({"Wide": [wide_row, ["v"] * (MAX_COLS_PER_SHEET + 10)]})
    storage = _FakeStorage({"test/file.xlsx": source})
    doc = _doc()

    SpreadsheetProvider().prepare(
        doc, storage_get=storage.get_sync, storage_put=storage.put_sync
    )

    sheet = json.loads(storage.store[f"viewable/{doc.id}.json"].decode("utf-8"))[
        "sheets"
    ][0]
    assert len(sheet["headers"]) == MAX_COLS_PER_SHEET
    assert sheet["total_cols"] == MAX_COLS_PER_SHEET + 10
    assert sheet["truncated"] is True


@pytest.mark.asyncio
async def test_render_without_viewable_key_is_conversion_pending() -> None:
    provider = SpreadsheetProvider()
    doc = _doc()

    class _AsyncStorage:
        async def get(self, key: str) -> bytes:
            raise AssertionError("render must not fetch when viewable_key is None")

        async def presigned_url(self, key: str, expires_seconds: int = 3600) -> str:
            raise AssertionError("not used for table kind")

    payload = await provider.render(doc, _AsyncStorage())
    assert payload.kind == "unsupported"
    assert payload.meta["reason"] == "conversion_pending"


@pytest.mark.asyncio
async def test_render_with_viewable_key_returns_inline_table() -> None:
    source = _xlsx({"Sheet1": [["h1", "h2"], ["v1", "v2"]]})
    provider = SpreadsheetProvider()
    doc = _doc()
    doc.viewable_key = f"viewable/{doc.id}.json"
    storage_sync = _FakeStorage({"test/file.xlsx": source})
    # Run prepare to get the cached JSON into the fake store.
    SpreadsheetProvider().prepare(
        doc,
        storage_get=storage_sync.get_sync,
        storage_put=storage_sync.put_sync,
    )

    class _AsyncStorage:
        def __init__(self, store):
            self._store = store

        async def get(self, key: str) -> bytes:
            return self._store[key]

        async def presigned_url(self, *args, **kwargs):
            raise AssertionError("tables use inline data, not signed URLs")

    payload = await provider.render(doc, _AsyncStorage(storage_sync.store))
    assert payload.kind == "table"
    assert payload.data is not None
    assert payload.data["sheets"][0]["headers"] == ["h1", "h2"]
    assert payload.meta["sheet_count"] == 1
